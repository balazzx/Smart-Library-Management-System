import os
import shutil
import openpyxl
from datetime import datetime
from app import create_app, db
from app.models.user import User, Department, Student, Librarian
from app.models.book import Author, Category, Publisher, BookLocation, Book, BookCopy
from app.models.transaction import BorrowTransaction, ReturnTransaction, BookRequest, Reservation, WaitingList
from app.models.system import Fine, SystemSetting, AuditLog
from app.models.engagement import Announcement

# ─── AUTHORITATIVE FILENAMES (Final verified datasets only) ───────────────────
STUDENTS_FILE     = 'Students_Final.xlsx'
BOOKS_FILE        = 'Books_Final.xlsx'
COPIES_FILE       = 'Book_Copies_Final.xlsx'
HISTORY_FILE      = 'Borrow_Return_History_Final.xlsx'
LOCATIONS_FILE    = 'Library_Locations_Final.xlsx'

def parse_date(val):
    if not val:
        return None
    if isinstance(val, datetime):
        return val
    val_str = str(val).strip()
    if not val_str or val_str.lower() == 'none':
        return None
    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d-%m-%Y', '%m/%d/%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(val_str, fmt)
        except ValueError:
            pass
    return None

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    return '' if s.lower() == 'none' else s

def inspect_headers(filepath, sheet_name=None):
    """Helper to inspect Excel headers without crashing."""
    try:
        wb = openpyxl.load_workbook(filepath)
        sheet = wb.active if not sheet_name else wb[sheet_name]
        headers = [clean_str(c.value) for c in sheet[1]]
        print(f"  Headers in {os.path.basename(filepath)}: {headers}")
        # Print first data row
        for row in sheet.iter_rows(min_row=2, max_row=2, values_only=True):
            print(f"  Sample row: {row}")
        return headers
    except Exception as e:
        print(f"  Could not inspect {filepath}: {e}")
        return []


def run_import():
    app = create_app()
    with app.app_context():

        datas_dir = os.path.join(app.root_path, '..', 'Datas')
        datas_dir = os.path.abspath(datas_dir)

        # ── VERIFY all required Final files exist ────────────────────────────
        required_files = [STUDENTS_FILE, BOOKS_FILE, COPIES_FILE, HISTORY_FILE, LOCATIONS_FILE]
        print("\n=== VERIFYING AUTHORITATIVE DATA FILES ===")
        all_ok = True
        for fname in required_files:
            fpath = os.path.join(datas_dir, fname)
            exists = os.path.exists(fpath)
            size = os.path.getsize(fpath) if exists else 0
            status = f"FOUND ({size} bytes)" if exists else "MISSING ❌"
            print(f"  {fname}: {status}")
            if not exists:
                all_ok = False
        if not all_ok:
            print("ERROR: One or more required Final files are missing. Aborting.")
            return

        # ── BACKUP existing database ─────────────────────────────────────────
        db_path = os.path.join(app.root_path, '..', 'smart_library.db')
        db_path = os.path.abspath(db_path)
        if os.path.exists(db_path):
            backup_path = db_path.replace('.db', f'_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.db')
            shutil.copy2(db_path, backup_path)
            print(f"\nDatabase backed up to: {os.path.basename(backup_path)}")

        print("\nRecreating database tables for fresh KIOT data import...")
        db.drop_all()
        db.create_all()

        # ── INSPECT Excel headers for key files ─────────────────────────────
        print("\n=== INSPECTING EXCEL FILE HEADERS ===")
        students_path  = os.path.join(datas_dir, STUDENTS_FILE)
        books_path     = os.path.join(datas_dir, BOOKS_FILE)
        copies_path    = os.path.join(datas_dir, COPIES_FILE)
        history_path   = os.path.join(datas_dir, HISTORY_FILE)
        locations_path = os.path.join(datas_dir, LOCATIONS_FILE)

        student_raw_headers = inspect_headers(students_path)
        books_raw_headers   = inspect_headers(books_path)
        copies_raw_headers  = inspect_headers(copies_path)
        history_raw_headers = inspect_headers(history_path)
        loc_raw_headers     = inspect_headers(locations_path)

        # ── 1. DEPARTMENTS ────────────────────────────────────────────────────
        print("\n=== 1. INITIALIZING KIOT DEPARTMENTS ===")
        dept_info = [
            ('CSE',       'Computer Science and Engineering'),
            ('IT',        'Information Technology'),
            ('AI & DS',   'Artificial Intelligence and Data Science'),
            ('CSBS',      'Computer Science and Business Systems'),
            ('ECE',       'Electronics and Communication Engineering'),
            ('EEE',       'Electrical and Electronics Engineering'),
            ('ECX',       'Electronics and Computer Engineering'),
            ('Mechanical','Mechanical Engineering'),
            ('Civil',     'Civil Engineering'),
            ('General',   'General Studies'),
        ]
        dept_map = {}
        for code, name in dept_info:
            dept = Department(code=code, name=name)
            db.session.add(dept)
        db.session.commit()

        all_depts = Department.query.all()
        for d in all_depts:
            dept_map[d.code.upper()]  = d.id
            dept_map[d.name.upper()]  = d.id

        # Aliases for common variations
        dept_map['AIANDS']  = dept_map.get('AI & DS')
        dept_map['AI&DS']   = dept_map.get('AI & DS')
        dept_map['AIDS']    = dept_map.get('AI & DS')
        dept_map['MECH']    = dept_map.get('MECHANICAL ENGINEERING')
        dept_map['CIVIL']   = dept_map.get('CIVIL ENGINEERING')
        if not dept_map.get('CIVIL'):
            dept_map['CIVIL'] = dept_map.get('CIVIL ENGINEERING') or dept_map.get('Civil'.upper())
        print(f"  Created {len(all_depts)} departments.")

        # ── 2. STAFF / ADMIN ACCOUNTS ─────────────────────────────────────────
        print("\n=== 2. CREATING ADMIN & LIBRARIAN ACCOUNTS ===")
        admin_user = User(username='kiot.admin', email='admin@kiot.ac.in', role='ADMIN')
        admin_user.set_password('KIOT@Admin2024')
        db.session.add(admin_user)

        lib_user = User(username='kiot.librarian', email='librarian@kiot.ac.in', role='LIBRARIAN')
        lib_user.set_password('KIOT@Library2024')
        db.session.add(lib_user)
        db.session.commit()

        cse_dept_id = dept_map.get('CSE')
        librarian_profile = Librarian(
            user_id=lib_user.id,
            staff_id='LIB-001',
            name='KIOT Chief Librarian',
            department_id=cse_dept_id,
            phone='9876543210'
        )
        db.session.add(librarian_profile)
        db.session.commit()
        print("  Admin: kiot.admin / KIOT@Admin2024")
        print("  Librarian: kiot.librarian / KIOT@Library2024")

        # ── 3. IMPORT LOCATIONS ───────────────────────────────────────────────
        print(f"\n=== 3. IMPORTING LOCATIONS FROM {LOCATIONS_FILE} ===")
        wb = openpyxl.load_workbook(locations_path)
        sheet = wb.active
        headers = [clean_str(cell.value) for cell in sheet[1]]

        location_map = {}
        loc_count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = dict(zip(headers, row))
            block    = clean_str(row_dict.get('Block'))
            floor    = clean_str(row_dict.get('Floor'))
            section  = clean_str(row_dict.get('Section'))
            rack     = clean_str(row_dict.get('Rack'))
            shelf    = clean_str(row_dict.get('Shelf'))
            position = clean_str(row_dict.get('Position'))

            loc = BookLocation(
                block=block, floor=floor, section=section,
                rack=rack, shelf=shelf, position=position
            )
            db.session.add(loc)
            db.session.flush()

            loc_key = (block, floor, section, rack, shelf, position)
            location_map[loc_key] = loc.id
            loc_count += 1

        db.session.commit()
        print(f"  Imported {loc_count} library locations.")

        # ── 4. IMPORT BOOKS ───────────────────────────────────────────────────
        print(f"\n=== 4. IMPORTING BOOKS FROM {BOOKS_FILE} ===")
        wb = openpyxl.load_workbook(books_path)
        sheet = wb.active
        headers = [clean_str(cell.value) for cell in sheet[1]]

        author_cache  = {}
        cat_cache     = {}
        pub_cache     = {}
        book_isbn_map = {}
        book_access_map = {}
        book_count = 0
        skipped_books = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = dict(zip(headers, row))
            isbn         = clean_str(row_dict.get('ISBN'))
            access_no    = clean_str(row_dict.get('ACCESS NO.'))
            title        = clean_str(row_dict.get('Book Title')) or clean_str(row_dict.get('Title'))
            author_name  = clean_str(row_dict.get('Author')) or 'Unknown Author'
            publisher_name = clean_str(row_dict.get('Publisher')) or 'KIOT Press'
            category_name  = clean_str(row_dict.get('Category')) or 'General'
            dept_name    = clean_str(row_dict.get('Department'))
            edition      = clean_str(row_dict.get('Edition'))
            subject      = clean_str(row_dict.get('Subject'))
            language     = clean_str(row_dict.get('Language')) or 'English'
            desc         = clean_str(row_dict.get('Description'))
            pub_year_val = row_dict.get('Publication Year')
            try:
                pub_year = int(pub_year_val)
            except (ValueError, TypeError):
                pub_year = 2024

            if not isbn and not title:
                continue

            # Guard against duplicate ISBN
            if isbn and isbn in book_isbn_map:
                skipped_books += 1
                continue

            # Author
            if author_name not in author_cache:
                auth = Author.query.filter_by(name=author_name).first()
                if not auth:
                    auth = Author(name=author_name)
                    db.session.add(auth)
                    db.session.flush()
                author_cache[author_name] = auth.id

            # Category
            if category_name not in cat_cache:
                cat = Category.query.filter_by(name=category_name).first()
                if not cat:
                    cat = Category(name=category_name)
                    db.session.add(cat)
                    db.session.flush()
                cat_cache[category_name] = cat.id

            # Publisher
            if publisher_name not in pub_cache:
                pub = Publisher.query.filter_by(name=publisher_name).first()
                if not pub:
                    pub = Publisher(name=publisher_name)
                    db.session.add(pub)
                    db.session.flush()
                pub_cache[publisher_name] = pub.id

            dept_id = dept_map.get(dept_name.upper()) if dept_name else None

            book = Book(
                title=title,
                isbn=isbn or f'ISBN-MISSING-{book_count+1}',
                access_no=access_no,
                author_id=author_cache[author_name],
                category_id=cat_cache[category_name],
                publisher_id=pub_cache[publisher_name],
                pub_year=pub_year,
                edition=edition,
                subject=subject,
                language=language,
                description=desc,
                department_id=dept_id
            )
            db.session.add(book)
            db.session.flush()

            if isbn:
                book_isbn_map[isbn] = book.id
            if access_no:
                book_access_map[access_no] = book.id
            book_count += 1

        db.session.commit()
        print(f"  Imported {book_count} books. Skipped {skipped_books} duplicates.")

        # ── 5. IMPORT BOOK COPIES ─────────────────────────────────────────────
        print(f"\n=== 5. IMPORTING BOOK COPIES FROM {COPIES_FILE} ===")
        wb = openpyxl.load_workbook(copies_path)
        sheet = wb.active
        headers = [clean_str(cell.value) for cell in sheet[1]]

        copy_code_map = {}
        copy_count = 0
        copy_skipped = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = dict(zip(headers, row))
            access_no = clean_str(row_dict.get('ACCESS NO.'))
            isbn      = clean_str(row_dict.get('ISBN'))
            cond_raw  = clean_str(row_dict.get('Condition'))
            cond      = cond_raw.upper() if cond_raw else 'EXCELLENT'
            status_raw = clean_str(row_dict.get('Status'))
            status    = status_raw.upper() if status_raw else 'AVAILABLE'

            block    = clean_str(row_dict.get('Block'))
            floor    = clean_str(row_dict.get('Floor'))
            section  = clean_str(row_dict.get('Section'))
            rack     = clean_str(row_dict.get('Rack'))
            shelf    = clean_str(row_dict.get('Shelf'))
            position = clean_str(row_dict.get('Position'))

            loc_id  = location_map.get((block, floor, section, rack, shelf, position))
            book_id = book_access_map.get(access_no) or book_isbn_map.get(isbn)

            if not book_id:
                print(f"  Warning: No book found for copy {access_no} (ISBN: {isbn})")
                copy_skipped += 1
                continue

            # Guard against duplicate access_no
            if access_no and access_no in copy_code_map:
                copy_skipped += 1
                continue

            copy = BookCopy(
                copy_code=access_no or f'COPY-MISSING-{copy_count+1}',
                book_id=book_id,
                condition_status=cond,
                status=status,
                location_id=loc_id
            )
            db.session.add(copy)
            db.session.flush()

            if access_no:
                copy_code_map[access_no] = copy.id
            copy_count += 1

        db.session.commit()
        print(f"  Imported {copy_count} book copies. Skipped {copy_skipped}.")

        # ── PRE-READ HISTORY for student department inference ─────────────────
        print(f"\n=== PRE-READING BORROW HISTORY FOR DEPT INFERENCE ===")
        student_dept_mapping = {}
        try:
            wb_h = openpyxl.load_workbook(history_path)
            s_h  = wb_h.active
            hist_headers = [clean_str(c.value) for c in s_h[1]]
            print(f"  History headers: {hist_headers}")

            # Try to find Roll No and Department columns by name
            roll_col = None
            dept_col = None
            for i, h in enumerate(hist_headers):
                hl = h.lower()
                if 'roll' in hl:
                    roll_col = i
                if 'dept' in hl or 'department' in hl:
                    dept_col = i

            if roll_col is not None and dept_col is not None:
                for r_h in s_h.iter_rows(min_row=2, values_only=True):
                    if r_h and r_h[roll_col] and r_h[dept_col]:
                        r_no = clean_str(r_h[roll_col]).upper()
                        d_code = clean_str(r_h[dept_col]).upper()
                        student_dept_mapping[r_no] = d_code
                print(f"  Inferred department for {len(student_dept_mapping)} students from history.")
            else:
                print(f"  WARNING: Could not find Roll No ({roll_col}) or Dept ({dept_col}) column in history. Student department will fallback to Excel data.")
        except Exception as e:
            print(f"  Warning: Could not pre-read history for dept inference: {e}")

        # ── 6. IMPORT STUDENTS ────────────────────────────────────────────────
        print(f"\n=== 6. IMPORTING STUDENTS FROM {STUDENTS_FILE} ===")
        wb    = openpyxl.load_workbook(students_path)
        sheet = wb.active

        # Normalize headers with flexible column mapping
        raw_headers = [clean_str(cell.value) for cell in sheet[1]]
        clean_headers = []
        for h in raw_headers:
            h_clean = h.rstrip(':').strip()
            h_low   = h_clean.lower()
            if 'roll' in h_low:
                clean_headers.append('roll_no')
            elif 'reg' in h_low:
                clean_headers.append('reg_no')
            elif 'dept' in h_low or 'department' in h_low:
                clean_headers.append('department')
            elif 'name' in h_low and 'mentor' not in h_low and 'advisor' not in h_low:
                clean_headers.append('name')
            elif 'year' in h_low:
                clean_headers.append('year')
            elif 'section' in h_low:
                clean_headers.append('section')
            elif 'gender' in h_low:
                clean_headers.append('gender')
            elif 'mail' in h_low or 'email' in h_low:
                clean_headers.append('email')
            elif 'mentor' in h_low:
                clean_headers.append('mentor_name')
            elif 'advisor' in h_low or 'class' in h_low:
                clean_headers.append('class_advisor')
            else:
                clean_headers.append(h_clean)

        print(f"  Mapped student headers: {clean_headers}")

        student_roll_map = {}
        student_count = 0
        skipped_students = 0
        dept_distribution = {}
        seen_emails = {u.email.lower() for u in User.query.all()}
        seen_rolls  = set()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = dict(zip(clean_headers, row))
            roll_no      = clean_str(row_dict.get('roll_no'))
            reg_no       = clean_str(row_dict.get('reg_no'))
            name         = clean_str(row_dict.get('name'))
            dept_in_excel= clean_str(row_dict.get('department'))
            year_val     = clean_str(row_dict.get('year', ''))
            section      = clean_str(row_dict.get('section'))
            gender       = clean_str(row_dict.get('gender'))
            raw_email    = clean_str(row_dict.get('email'))
            mentor       = clean_str(row_dict.get('mentor_name'))
            advisor      = clean_str(row_dict.get('class_advisor'))

            if not roll_no:
                skipped_students += 1
                continue

            # Skip duplicate roll numbers
            roll_upper = roll_no.upper()
            if roll_upper in seen_rolls:
                print(f"  Warning: Duplicate roll number {roll_no}, skipping.")
                skipped_students += 1
                continue
            seen_rolls.add(roll_upper)

            # Determine department: Priority: Excel column > history inference > prefix-based
            dept_code = None
            if dept_in_excel:
                dept_code = dept_in_excel.upper()
            if not dept_code:
                dept_code = student_dept_mapping.get(roll_upper)
            # Fallback: infer from roll number prefix (e.g. 2k24ecx023 → ECX)
            if not dept_code:
                roll_lower = roll_no.lower()
                if '2k24cse' in roll_lower:
                    dept_code = 'CSE'
                elif '2k24it' in roll_lower:
                    dept_code = 'IT'
                elif '2k24aids' in roll_lower or '2k24aiands' in roll_lower or '2k24ai' in roll_lower:
                    dept_code = 'AI & DS'
                elif '2k24csbs' in roll_lower:
                    dept_code = 'CSBS'
                elif '2k24ece' in roll_lower:
                    dept_code = 'ECE'
                elif '2k24eee' in roll_lower:
                    dept_code = 'EEE'
                elif '2k24ecx' in roll_lower:
                    dept_code = 'ECX'
                elif '2k24mech' in roll_lower:
                    dept_code = 'Mechanical'.upper()
                elif '2k24civil' in roll_lower:
                    dept_code = 'Civil'.upper()

            dept_id = None
            if dept_code:
                dept_id = dept_map.get(dept_code)
            if not dept_id:
                dept_id = dept_map.get('CSE')  # Last resort fallback

            # Resolve human-readable code for display
            dept_record = Department.query.get(dept_id) if dept_id else None
            actual_dept_code = dept_record.code if dept_record else 'CSE'
            dept_distribution[actual_dept_code] = dept_distribution.get(actual_dept_code, 0) + 1

            # Email deduplication
            email = raw_email if (raw_email and raw_email.lower() not in seen_emails) else f"{roll_no.lower()}@kiot.ac.in"
            if email.lower() in seen_emails:
                email = f"{roll_no.lower()}.std@kiot.ac.in"
            seen_emails.add(email.lower())

            # Year as integer
            try:
                yr_int = int(float(year_val)) if year_val else 1
            except (ValueError, TypeError):
                yr_int = 1

            username = roll_no.lower()
            u = User(username=username, email=email, role='STUDENT')
            u.set_password(username)  # Default password = roll number
            db.session.add(u)
            db.session.flush()

            card_no = f"LIB-STD-{roll_no.upper()}"
            st = Student(
                user_id=u.id,
                student_id=roll_no,
                reg_no=reg_no,
                name=name,
                department_id=dept_id,
                section=section,
                gender=gender,
                mentor_name=mentor,
                class_advisor=advisor,
                library_card_number=card_no
            )
            db.session.add(st)
            db.session.flush()

            student_roll_map[roll_upper] = st.id
            student_count += 1

        db.session.commit()
        print(f"  Imported {student_count} students. Skipped {skipped_students}.")
        print(f"  Department distribution: {dept_distribution}")

        # ── 7. IMPORT BORROW & RETURN HISTORY ────────────────────────────────
        print(f"\n=== 7. IMPORTING BORROW/RETURN HISTORY FROM {HISTORY_FILE} ===")
        wb    = openpyxl.load_workbook(history_path)
        sheet = wb.active
        headers = [clean_str(cell.value) for cell in sheet[1]]
        print(f"  Transaction headers: {headers}")

        tx_count   = 0
        fine_count = 0
        tx_skipped = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = dict(zip(headers, row))

            tx_ref      = clean_str(row_dict.get('Transaction ID'))
            roll_no     = clean_str(row_dict.get('Student Roll No')).upper()
            access_no   = clean_str(row_dict.get('Book Access No'))
            issue_date  = parse_date(row_dict.get('Issue Date'))
            due_date    = parse_date(row_dict.get('Due Date'))
            return_date = parse_date(row_dict.get('Return Date'))
            status_raw  = clean_str(row_dict.get('Status'))
            status      = status_raw.upper() if status_raw else 'ISSUED'

            try:
                renewal_cnt = int(row_dict.get('Renewal Count') or 0)
            except (ValueError, TypeError):
                renewal_cnt = 0

            try:
                fine_amt = float(row_dict.get('Fine Amount') or 0.0)
            except (ValueError, TypeError):
                fine_amt = 0.0

            issued_by   = clean_str(row_dict.get('Issued By'))
            returned_to = clean_str(row_dict.get('Returned To'))
            remarks     = clean_str(row_dict.get('Remarks'))

            st_id   = student_roll_map.get(roll_no)
            copy_id = copy_code_map.get(access_no)

            if not st_id or not copy_id:
                tx_skipped += 1
                if tx_skipped <= 20:
                    print(f"  Skip txn {tx_ref}: Roll={roll_no}(found={st_id is not None}), Copy={access_no}(found={copy_id is not None})")
                continue

            # Skip duplicate transaction refs
            if tx_ref and BorrowTransaction.query.filter_by(transaction_ref=tx_ref).first():
                tx_skipped += 1
                continue

            tx = BorrowTransaction(
                transaction_ref=tx_ref or None,
                copy_id=copy_id,
                student_id=st_id,
                issue_date=issue_date or datetime.utcnow(),
                due_date=due_date or datetime.utcnow(),
                return_date=return_date,
                status=status,
                fine_amount=fine_amt,
                is_fine_paid=(status == 'RETURNED' and fine_amt > 0),
                renewal_count=renewal_cnt,
                issued_by=issued_by,
                returned_to=returned_to,
                remarks=remarks
            )
            db.session.add(tx)
            db.session.flush()

            # Mark copy as ISSUED if still active
            if status in ('ISSUED', 'OVERDUE'):
                c = db.session.get(BookCopy, copy_id)
                if c:
                    c.status = 'ISSUED'

            # Create Fine record if fine > 0
            if fine_amt > 0:
                fine = Fine(
                    student_id=st_id,
                    transaction_id=tx.id,
                    amount=fine_amt,
                    reason=f"Late return fee for access no {access_no}",
                    status=('PAID' if status == 'RETURNED' else 'UNPAID'),
                    paid_at=(return_date if status == 'RETURNED' else None)
                )
                db.session.add(fine)
                fine_count += 1

            tx_count += 1

        db.session.commit()
        print(f"  Imported {tx_count} borrow transactions, {fine_count} fines.")
        print(f"  Skipped {tx_skipped} transactions (unmatched roll no / access no).")

        # ── 8. SYSTEM SETTINGS & ANNOUNCEMENTS ───────────────────────────────
        print("\n=== 8. CONFIGURING SYSTEM SETTINGS ===")
        settings_data = [
            ('LIBRARY_NAME',         'KIOT CENTRAL LIBRARY',          'Official library name'),
            ('FINE_RATE_PER_DAY',     '2.0',                           'Daily fine rate in ₹'),
            ('MAX_BORROW_DAYS',       '14',                            'Maximum borrow period in days'),
            ('MAX_BORROW_LIMIT',      '5',                             'Maximum books a student can borrow'),
            ('BORROW_PERIOD_DAYS',    '14',                            'Borrow period in days'),
            ('GRACE_PERIOD_DAYS',     '2',                             'Grace period in days before fines'),
            ('FINE_PER_DAY',          '2.0',                           'Fine per day in ₹'),
            ('MAX_FINE',              '500.0',                         'Maximum fine amount in ₹'),
            ('CLAIM_PERIOD_HOURS',    '48',                            'Hours to claim reserved book'),
            ('LIBRARY_EMAIL',         'library@kiot.ac.in',            'Library contact email'),
            ('LIBRARY_PHONE',         '+91 427 2452371',               'Library phone number'),
            ('COLLEGE_NAME',          'Knowledge Institute of Technology', 'College full name'),
            ('COLLEGE_ADDRESS',       'NH-544, Kakapalayam, Salem, Tamil Nadu - 637504', 'College address'),
        ]
        for key, val, desc in settings_data:
            s = SystemSetting.query.filter_by(setting_key=key).first()
            if not s:
                s = SystemSetting(setting_key=key, setting_value=val, description=desc)
                db.session.add(s)
            else:
                s.setting_value = val

        today = datetime.utcnow().date()
        welcome_ann = Announcement(
            title="Welcome to KIOT Central Library Portal!",
            message=(
                "The Knowledge Institute of Technology (KIOT) Central Library portal is now fully operational "
                "with real catalog, student, and circulation records. Students can search the catalog, manage borrowings, "
                "view digital resources, and track their library activity. For assistance, contact the library at librarian@kiot.ac.in."
            ),
            priority='HIGH',
            target_department_id=None,
            target_role='ALL',
            start_date=today,
            end_date=datetime(today.year + 2, 12, 31).date(),
            created_by=admin_user.id
        )
        db.session.add(welcome_ann)
        db.session.commit()

        # ── COPY LOGO to static if not already present ────────────────────────
        logo_src = os.path.join(datas_dir, 'kiot_logo.jpg')
        logo_dst = os.path.join(app.root_path, 'static', 'uploads', 'kiot_logo.jpg')
        if os.path.exists(logo_src) and not os.path.exists(logo_dst):
            os.makedirs(os.path.dirname(logo_dst), exist_ok=True)
            shutil.copy2(logo_src, logo_dst)
            print(f"\nKIOT logo copied to static/uploads/kiot_logo.jpg")
        elif not os.path.exists(logo_dst):
            print("WARNING: kiot_logo.jpg not found in Datas/")

        # ── FINAL SUMMARY ─────────────────────────────────────────────────────
        print("\n" + "="*60)
        print("=== KIOT CENTRAL LIBRARY — IMPORT COMPLETE ===")
        print("="*60)
        print(f"  Departments       : {Department.query.count()}")
        print(f"  Users             : {User.query.count()}")
        print(f"  Students          : {Student.query.count()}")
        print(f"  Locations         : {BookLocation.query.count()}")
        print(f"  Books             : {Book.query.count()}")
        print(f"  Book Copies       : {BookCopy.query.count()}")
        print(f"  Borrow Transactions: {BorrowTransaction.query.count()}")
        print(f"  Fines             : {Fine.query.count()}")
        print()
        print("=== DEPARTMENT DISTRIBUTION ===")
        for dept in Department.query.all():
            cnt = Student.query.filter_by(department_id=dept.id).count()
            if cnt > 0:
                print(f"  {dept.code:15s}: {cnt}")
        print()

        # Verify expected counts
        expected = {'Students': 66, 'Books': 500, 'Book Copies': 500, 'Locations': 360}
        actual   = {
            'Students': Student.query.count(),
            'Books': Book.query.count(),
            'Book Copies': BookCopy.query.count(),
            'Locations': BookLocation.query.count(),
        }
        print("=== VERIFICATION ===")
        for k, expected_val in expected.items():
            actual_val = actual[k]
            ok = "PASS" if actual_val == expected_val else f"WARNING (expected {expected_val}, got {actual_val})"
            print(f"  {k:20s}: {actual_val:4d}  {ok}")

        print()
        print("Credentials:")
        print("  Admin       → kiot.admin     / KIOT@Admin2024")
        print("  Librarian   → kiot.librarian / KIOT@Library2024")
        print("  Students    → (roll_no) / (roll_no)")
        print("="*60)

if __name__ == '__main__':
    run_import()
